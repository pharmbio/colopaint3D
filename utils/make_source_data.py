#!/usr/bin/env python3
"""Build the journal Source Data workbook from the per-panel tables.

Nature's requirement: "the relevant raw data from each figure or table ... should be
represented by a single sheet in an Excel document". This assembles exactly that from
``source_data/*.csv`` and ``source_data/MANIFEST.csv`` — one sheet per panel, in figure
order, preceded by a contents sheet naming every panel, its description and its row
count.

    python utils/make_source_data.py            # -> Source Data.xlsx
    python utils/make_source_data.py --out X.xlsx

Run ``python run_all.py`` first: a panel with no table is an error here, because a
Source Data file that quietly omits a panel is worse than one that fails to build.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
SOURCE_DATA = REPO / "source_data"
MANIFEST = SOURCE_DATA / "MANIFEST.csv"

# Excel caps sheet names at 31 characters and forbids : \ / ? * [ ]
_BAD_SHEET = re.compile(r"[:\\/?*\[\]]")


def _sheet_name(panel: str) -> str:
    return _BAD_SHEET.sub("-", panel)[:31]


def _panel_sort_key(panel: str) -> tuple:
    """Main figures before supplementary, then figure number, then panel letter."""
    m = re.match(r"^(Suppl)?Fig(\d+)([A-Za-z]?)(\d?)_?(.*)$", panel)
    if not m:
        return (2, 99, "z", "", panel)
    suppl, num, letter, sub, part = m.groups()
    return (1 if suppl else 0, int(num), letter, sub, part)


def build(out_path: Path) -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("needs openpyxl: pip install openpyxl", file=sys.stderr)
        return 2

    if not MANIFEST.exists():
        print(f"no {MANIFEST}. Run: python run_all.py", file=sys.stderr)
        return 1

    with MANIFEST.open(newline="") as fh:
        rows = sorted(csv.DictReader(fh), key=lambda r: _panel_sort_key(r["panel"]))
    if not rows:
        print("MANIFEST.csv is empty. Run: python run_all.py", file=sys.stderr)
        return 1

    missing = [r["panel"] for r in rows if not (SOURCE_DATA / r["source_data"]).exists()]
    if missing:
        print(f"tables missing for {len(missing)} panel(s): {', '.join(missing)}", file=sys.stderr)
        return 1

    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    contents.append(["Source Data"])
    contents["A1"].font = Font(bold=True, size=14)
    contents.append([
        "High-content morphological profiling by Cell Painting in 3D spheroids"
    ])
    contents.append([])
    contents.append(["Panel", "Figure", "Description", "Rows", "Produced by"])
    for cell in contents[4]:
        cell.font = Font(bold=True)

    for r in rows:
        contents.append([r["panel"], r["figure"], r["caption"],
                         int(r["n_rows"] or 0), r["notebook"]])

    for width, col in zip((22, 14, 74, 9, 58), "ABCDE"):
        contents.column_dimensions[col].width = width
    for row in contents.iter_rows(min_row=5):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    used: set[str] = set()
    for r in rows:
        name = _sheet_name(r["panel"])
        # panel names are unique in the manifest, but truncation to 31 chars could collide
        n = 1
        while name in used:
            n += 1
            name = f"{_sheet_name(r['panel'])[:29]}_{n}"
        used.add(name)

        ws = wb.create_sheet(name)
        with (SOURCE_DATA / r["source_data"]).open(newline="") as fh:
            for i, line in enumerate(csv.reader(fh)):
                ws.append([_maybe_number(v) for v in line])
                if i == 0:
                    for cell in ws[1]:
                        cell.font = Font(bold=True)
        ws.freeze_panes = "A2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"wrote {out_path}  —  {len(rows)} panel sheets + contents")
    return 0


def _maybe_number(v: str):
    """Keep numbers numeric so the sheets are usable, leave everything else as text."""
    if v in ("", "NA", "NaN", "nan"):
        return None
    try:
        return int(v)
    except ValueError:
        pass
    try:
        return float(v)
    except ValueError:
        return v


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(REPO / "Source Data.xlsx"),
                    help="output path (default: 'Source Data.xlsx' at the repo root)")
    return build(Path(ap.parse_args().out))


if __name__ == "__main__":
    raise SystemExit(main())
